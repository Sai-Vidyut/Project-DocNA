"""Run the full pipeline on real-world fixtures and collect validation reports."""

from __future__ import annotations

import hashlib
import json
import subprocess
import sys
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path

from docx import Document

from docna.adapters.docx import DocxAdapter
from docna.adapters.docx.locators import DOCX_MAIN_PART
from docna.answer.context import build_context_packs
from docna.answer.generate import generate_answers
from docna.config import Settings
from docna.detect.run import detect_from_adapter
from docna.ir import PlacementOp
from docna.place.planner import plan_placements
from docna.serializers import summarize_placement_ops
from tests.pipeline_helpers import pipeline_mock_provider
from tests.placement_helpers import apply_ops, sha256_file

REALWORLD_DIR = Path(__file__).resolve().parent / "fixtures" / "docx" / "realworld"
GENERATOR = REALWORLD_DIR.parent / "generate_realworld_fixtures.py"


@dataclass
class FixtureReport:
    fixture: str
    status: str
    original_hash: str
    output_hash: str | None = None
    original_unchanged: bool = False
    output_valid_zip: bool = False
    output_opens: bool = False
    table_count_before: int = 0
    table_count_after: int = 0
    answer_spaces: int = 0
    tasks_total: int = 0
    tasks_answerable: int = 0
    tasks_skipped: int = 0
    tasks_instruction: int = 0
    tasks_already_answered: int = 0
    answers_generated: int = 0
    placement_strategies: dict[str, int] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    detected_prompts: list[str] = field(default_factory=list)
    skipped_prompts: list[str] = field(default_factory=list)


def ensure_realworld_fixtures() -> list[Path]:
    if not REALWORLD_DIR.exists() or not any(REALWORLD_DIR.glob("*.docx")):
        subprocess.run([sys.executable, str(GENERATOR)], check=True)
    return sorted(REALWORLD_DIR.glob("*.docx"))


def run_fixture(fixture_path: Path, work_dir: Path) -> FixtureReport:
    before_hash = sha256_file(fixture_path)
    adapter = DocxAdapter()
    provider = pipeline_mock_provider()
    report = FixtureReport(fixture=fixture_path.name, status="pending", original_hash=before_hash)

    try:
        ir = adapter.parse(fixture_path)
        report.answer_spaces = len(ir.answer_spaces)
        report.table_count_before = len(ir.tables)
        report.warnings.extend(ir.warnings)

        detection = detect_from_adapter(ir, adapter, provider)
        report.tasks_total = len(detection.tasks)
        report.tasks_instruction = sum(1 for t in detection.tasks if t.kind == "instruction")
        report.tasks_already_answered = sum(
            1 for t in detection.tasks if t.skip_reason == "already_answered"
        )
        report.tasks_skipped = sum(
            1
            for t in detection.tasks
            if t.skip_reason and t.skip_reason != "already_answered"
        )
        report.tasks_answerable = len(detection.answerable_tasks)
        report.detected_prompts = [t.prompt_text[:80] for t in detection.tasks]
        report.skipped_prompts = [
            t.prompt_text[:80]
            for t in detection.tasks
            if t.skip_reason or t.kind == "instruction"
        ]
        report.warnings.extend(detection.warnings)

        packs = build_context_packs(ir, detection.tasks)
        generation = generate_answers(packs, provider)
        report.answers_generated = len(generation.answers)
        report.warnings.extend(generation.warnings)
        report.errors.extend(f"{e.task_id}:{e.error_type}" for e in generation.errors)

        ops = plan_placements(ir, detection.tasks, generation.answers)
        summaries = summarize_placement_ops(ops)
        strategy_counts: dict[str, int] = {}
        for op in summaries:
            strategy_counts[op.strategy] = strategy_counts.get(op.strategy, 0) + 1
        report.placement_strategies = strategy_counts

        output = apply_ops(fixture_path, ops, work_dir)
        report.output_hash = sha256_file(output)
        report.original_unchanged = sha256_file(fixture_path) == before_hash
        report.output_valid_zip = zipfile.is_zipfile(output)
        report.output_opens = _docx_opens(output)

        reparsed = adapter.parse(output)
        report.table_count_after = len(reparsed.tables)
        _validate_writes(ops, reparsed, report)
        report.status = "completed"
    except Exception as exc:  # noqa: BLE001
        report.status = "failed"
        report.errors.append(str(exc))
        report.original_unchanged = sha256_file(fixture_path) == before_hash

    return report


def _docx_opens(path: Path) -> bool:
    try:
        Document(path)
        with zipfile.ZipFile(path) as archive:
            archive.read(DOCX_MAIN_PART)
        return True
    except Exception:  # noqa: BLE001
        return False


def _validate_writes(ops: list[PlacementOp], reparsed, report: FixtureReport) -> None:
    from docna.pipeline import _answer_text_present

    block_text = "\n".join(block.text for block in reparsed.blocks)
    for op in ops:
        if op.strategy == "skip" or not op.text.strip():
            continue
        if not _answer_text_present(op, block_text):
            report.errors.append(f"missing_write:{op.task_id}")


def run_all(work_dir: Path) -> list[FixtureReport]:
    fixtures = ensure_realworld_fixtures()
    return [run_fixture(path, work_dir / path.stem) for path in fixtures]


def main() -> None:
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        reports = run_all(Path(tmp))
        summary = {
            "fixtures_tested": len(reports),
            "completed": sum(1 for r in reports if r.status == "completed"),
            "failed": sum(1 for r in reports if r.status == "failed"),
            "reports": [asdict(r) for r in reports],
        }
        print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
