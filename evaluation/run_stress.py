"""Phase 12 real-world stress-test runner.

Runs the stress corpus through mock and optional real-AI pipelines, editable
review checks, and DOCX structure inspection.

Usage:
    uv run python -m evaluation.run_stress
    uv run python -m evaluation.run_stress --real-ai
    uv run python -m evaluation.run_stress --fixture grant_proposal_messy.docx
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sys
import time
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import UTC, datetime
from pathlib import Path

from docx import Document

from docna.adapters.docx import DocxAdapter
from docna.adapters.docx.locators import DOCX_MAIN_PART
from docna.ai.factory import create_ai_provider
from docna.ai.instrumented import InstrumentedAIProvider
from docna.answer.context import build_context_packs
from docna.answer.generate import generate_answers
from docna.answer.generation_config import GenerationConfig
from docna.config import Settings
from docna.detect.run import detect_from_adapter
from docna.edit_apply import apply_edited_answers
from docna.ir import Answer, PlacementOp, Task
from docna.pipeline import _answer_text_present, run_job
from docna.place.planner import plan_placements
from docna.review_edits import AnswerEdit
from docna.serializers import summarize_placement_ops
from docna.storage.jobs import JobStore
from evaluation.run_eval import (
    EvalIssue,
    _check_answers,
    _check_placement,
    _compare_detection,
    load_expectations as load_realworld_expectations,
)
from tests.pipeline_helpers import pipeline_mock_provider
from tests.placement_helpers import apply_ops, sha256_file
from tests.stress_helpers import ensure_all_stress_fixtures, ensure_stress_fixture

EVAL_ROOT = Path(__file__).resolve().parent
STRESS_EXPECTATIONS_PATH = EVAL_ROOT / "stress_expectations.json"

LAYER_CLASS = {
    "parsing": "A",
    "detection": "B",
    "context": "C",
    "answer": "D",
    "placement": "E",
    "formatting": "F",
    "review": "G",
    "provider": "H",
    "harness": "I",
    "pipeline": "E",
}


@dataclass
class StructureInspection:
    tables_before: int = 0
    tables_after: int = 0
    blocks_before: int = 0
    blocks_after: int = 0
    output_valid_zip: bool = False
    output_opens: bool = False
    original_hash: str = ""
    output_hash: str | None = None
    original_unchanged: bool = False
    notes: list[str] = field(default_factory=list)


@dataclass
class EditInspection:
    attempted: bool = False
    single_edit_ok: bool = False
    multi_edit_ok: bool = False
    original_unchanged: bool = False
    ai_output_preserved: bool = False
    edited_text_found: bool = False
    untouched_answer_preserved: bool = False
    errors: list[str] = field(default_factory=list)


@dataclass
class StressFixtureResult:
    fixture: str
    provider_mode: str
    status: str
    elapsed_ms: float
    detected_tasks: int = 0
    expected_min_tasks: int | None = None
    answerable_tasks: int = 0
    answers_generated: int = 0
    low_confidence_count: int = 0
    placement_strategies: dict[str, int] = field(default_factory=dict)
    missing_writes: list[str] = field(default_factory=list)
    structure: StructureInspection = field(default_factory=StructureInspection)
    edit: EditInspection = field(default_factory=EditInspection)
    issues: list[EvalIssue] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def load_stress_expectations() -> dict:
    if not STRESS_EXPECTATIONS_PATH.exists():
        raise FileNotFoundError(
            f"Missing {STRESS_EXPECTATIONS_PATH}. "
            "Run: uv run python -c \"from evaluation.generate_stress_expectations import main; main()\""
        )
    return json.loads(STRESS_EXPECTATIONS_PATH.read_text(encoding="utf-8"))


def _issue_counts(results: list[StressFixtureResult]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for result in results:
        for issue in result.issues:
            key = f"{issue.layer}:{issue.severity}"
            counts[key] = counts.get(key, 0) + 1
    return counts


def _classified_issues(results: list[StressFixtureResult]) -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = {letter: [] for letter in "ABCDEFGHI"}
    for result in results:
        for issue in result.issues:
            letter = LAYER_CLASS.get(issue.layer, "I")
            grouped[letter].append(
                {
                    "fixture": result.fixture,
                    "layer": issue.layer,
                    "severity": issue.severity,
                    "message": issue.message,
                    "task_id": issue.task_id,
                }
            )
    return grouped


def _inspect_structure(
    fixture_path: Path,
    output_path: Path | None,
    before_hash: str,
) -> StructureInspection:
    adapter = DocxAdapter()
    inspection = StructureInspection(original_hash=before_hash)
    ir_before = adapter.parse(fixture_path)
    inspection.tables_before = len(ir_before.tables)
    inspection.blocks_before = len(ir_before.blocks)
    inspection.notes.extend(ir_before.warnings)

    if output_path is None or not output_path.exists():
        inspection.notes.append("No output DOCX produced")
        return inspection

    inspection.output_valid_zip = zipfile.is_zipfile(output_path)
    try:
        Document(output_path)
        with zipfile.ZipFile(output_path) as archive:
            archive.read(DOCX_MAIN_PART)
        inspection.output_opens = True
    except Exception as exc:  # noqa: BLE001
        inspection.notes.append(f"Output open failed: {exc}")

    reparsed = adapter.parse(output_path)
    inspection.tables_after = len(reparsed.tables)
    inspection.blocks_after = len(reparsed.blocks)
    inspection.output_hash = sha256_file(output_path)
    inspection.original_unchanged = sha256_file(fixture_path) == before_hash

    if inspection.tables_before != inspection.tables_after:
        inspection.notes.append(
            f"Table count changed: {inspection.tables_before} -> {inspection.tables_after}"
        )
    return inspection


def _run_edit_checks(
    fixture_path: Path,
    *,
    storage_dir: Path,
    settings: Settings,
) -> EditInspection:
    inspection = EditInspection()
    if fixture_path.stat().st_size == 0:
        inspection.errors.append("Empty fixture")
        return inspection

    provider = pipeline_mock_provider()
    paths = JobStore(storage_dir).create_job(fixture_path.read_bytes(), fixture_path.name)
    result = run_job(JobStore(storage_dir), paths, provider, settings=settings)
    if result.status != "completed":
        inspection.errors.append(f"Pipeline status={result.status}")
        return inspection

    store = JobStore(storage_dir)
    paths = store.get_paths(paths.job_id)
    review = store.load_review(paths.job_id)
    editable = [item for item in review.get("answered", []) if item.get("editable")]
    if len(editable) < 2:
        inspection.errors.append("Not enough editable answers for multi-edit test")
        return inspection

    inspection.attempted = True
    before_original = sha256_file(paths.original)
    ai_hash = sha256_file(paths.output_file)

    first = editable[0]
    second = editable[1]
    single_text = "Stress single edited answer."
    multi_text = "Stress multi edited answer."

    try:
        apply_edited_answers(
            store,
            paths,
            edits=[AnswerEdit(task_id=first["task_id"], text=single_text)],
        )
        inspection.single_edit_ok = True
    except Exception as exc:  # noqa: BLE001
        inspection.errors.append(f"Single edit failed: {exc}")
        return inspection

    reparsed = DocxAdapter().parse(paths.edited_output_file)
    block_text = "\n".join(block.text for block in reparsed.blocks)
    inspection.edited_text_found = single_text in block_text

    try:
        apply_edited_answers(
            store,
            paths,
            edits=[
                AnswerEdit(task_id=first["task_id"], text=single_text),
                AnswerEdit(task_id=second["task_id"], text=multi_text),
            ],
        )
        inspection.multi_edit_ok = True
        reparsed = DocxAdapter().parse(paths.edited_output_file)
        block_text = "\n".join(block.text for block in reparsed.blocks)
        inspection.edited_text_found = multi_text in block_text
        inspection.untouched_answer_preserved = any(
            item["answer_text"] in block_text
            for item in editable[2:5]
            if item.get("answer_text")
        )
    except Exception as exc:  # noqa: BLE001
        inspection.errors.append(f"Multi edit failed: {exc}")

    inspection.original_unchanged = sha256_file(paths.original) == before_original
    inspection.ai_output_preserved = sha256_file(paths.output_file) == ai_hash
    if not inspection.edited_text_found:
        inspection.errors.append("Edited text not found in final DOCX")
    return inspection


def evaluate_stress_fixture(
    fixture_path: Path,
    *,
    expectations: dict,
    provider,
    provider_mode: str,
    settings: Settings,
    work_dir: Path,
    storage_dir: Path,
    run_edit_checks: bool,
) -> StressFixtureResult:
    started = time.perf_counter()
    adapter = DocxAdapter()
    issues: list[EvalIssue] = []
    warnings: list[str] = []
    before_hash = sha256_file(fixture_path)
    spec = expectations.get(fixture_path.name, {})

    ir = adapter.parse(fixture_path)
    warnings.extend(ir.warnings)

    mock_detection = detect_from_adapter(ir, adapter, pipeline_mock_provider())
    detection = detect_from_adapter(ir, adapter, provider)
    tasks = detection.tasks
    warnings.extend(detection.warnings)
    issues.extend(_compare_detection(fixture_path.name, tasks, expectations, mock_detection.tasks))

    packs = build_context_packs(ir, tasks)
    generation = generate_answers(
        packs,
        provider,
        config=GenerationConfig(
            max_concurrency=settings.max_concurrency,
            batch_size=settings.answer_batch_size,
        ),
    )
    answers = generation.answers
    warnings.extend(generation.warnings)
    for error in generation.errors:
        issues.append(
            EvalIssue(
                layer="answer",
                severity="error",
                message=f"{error.error_type}: {error.message}",
                task_id=error.task_id,
            )
        )
    issues.extend(_check_answers(answers, tasks))

    ops = plan_placements(ir, tasks, answers)
    summaries = summarize_placement_ops(ops)
    strategy_counts: dict[str, int] = {}
    for summary in summaries:
        strategy_counts[summary.strategy] = strategy_counts.get(summary.strategy, 0) + 1

    output_path = None
    status = "completed"
    missing_writes: list[str] = []
    try:
        fixture_work = work_dir / fixture_path.stem
        fixture_work.mkdir(parents=True, exist_ok=True)
        output_path = apply_ops(fixture_path, ops, fixture_work)
        reparsed = adapter.parse(output_path)
        issues.extend(_check_placement(ops, reparsed, answers))
        block_text = "\n".join(block.text for block in reparsed.blocks)
        for op in ops:
            if op.strategy == "skip" or not op.text.strip():
                continue
            if not _answer_text_present(op, block_text):
                missing_writes.append(op.task_id)
                issues.append(
                    EvalIssue(
                        layer="placement",
                        severity="error",
                        message=f"Missing write for strategy {op.strategy}",
                        task_id=op.task_id,
                    )
                )
    except Exception as exc:  # noqa: BLE001
        status = "failed"
        issues.append(EvalIssue(layer="pipeline", severity="error", message=str(exc)))

    structure = _inspect_structure(fixture_path, output_path, before_hash)
    if not structure.output_opens and output_path is not None:
        issues.append(
            EvalIssue(layer="formatting", severity="error", message="Output DOCX failed to open")
        )
    if not structure.original_unchanged:
        issues.append(
            EvalIssue(layer="formatting", severity="error", message="Original DOCX hash changed")
        )

    edit = EditInspection()
    if run_edit_checks and provider_mode == "mock":
        edit = _run_edit_checks(fixture_path, storage_dir=storage_dir / fixture_path.stem, settings=settings)
        if edit.attempted and edit.errors:
            for message in edit.errors:
                issues.append(EvalIssue(layer="review", severity="error", message=message))
        if edit.attempted and not edit.original_unchanged:
            issues.append(
                EvalIssue(layer="review", severity="error", message="Original changed during edit apply")
            )

    low_confidence = sum(1 for answer in answers if answer.confidence < 0.70)
    if any(issue.severity == "error" for issue in issues):
        status = "failed"

    elapsed_ms = (time.perf_counter() - started) * 1000
    return StressFixtureResult(
        fixture=fixture_path.name,
        provider_mode=provider_mode,
        status=status,
        elapsed_ms=elapsed_ms,
        detected_tasks=len(tasks),
        expected_min_tasks=spec.get("min_answerable_tasks"),
        answerable_tasks=len(detection.answerable_tasks),
        answers_generated=len(answers),
        low_confidence_count=low_confidence,
        placement_strategies=strategy_counts,
        missing_writes=missing_writes,
        structure=structure,
        edit=edit,
        issues=issues,
        warnings=warnings,
    )


def _render_inspection_report(results: list[StressFixtureResult], run_dir: Path) -> str:
    lines = [
        "# DocNA Stress Test Inspection Report",
        "",
        f"Generated: {datetime.now(UTC).isoformat()}",
        f"Fixtures: {len(results)}",
        "",
        "## Summary",
        "",
    ]
    passed = sum(1 for result in results if result.status == "completed" and not any(i.severity == "error" for i in result.issues))
    failed = len(results) - passed
    lines.extend(
        [
            f"- Passed (no errors): {passed}",
            f"- Failed or issues: {failed}",
            "",
            "## Fixture Results",
            "",
        ]
    )
    for result in results:
        errors = [issue for issue in result.issues if issue.severity == "error"]
        lines.append(f"### {result.fixture}")
        lines.append("")
        lines.append(f"- Provider: {result.provider_mode}")
        lines.append(f"- Status: {result.status}")
        lines.append(f"- Detected tasks: {result.detected_tasks}")
        lines.append(f"- Answerable: {result.answerable_tasks}")
        lines.append(f"- Answers generated: {result.answers_generated}")
        lines.append(f"- Placement strategies: {result.placement_strategies}")
        lines.append(
            f"- Structure: tables {result.structure.tables_before}->{result.structure.tables_after}, "
            f"blocks {result.structure.blocks_before}->{result.structure.blocks_after}"
        )
        if result.edit.attempted:
            lines.append(
                f"- Edit checks: single={result.edit.single_edit_ok}, "
                f"multi={result.edit.multi_edit_ok}, ai_preserved={result.edit.ai_output_preserved}"
            )
        if errors:
            lines.append("- Errors:")
            for issue in errors[:8]:
                letter = LAYER_CLASS.get(issue.layer, "I")
                lines.append(f"  - [{letter}] {issue.message}")
        if result.structure.notes:
            lines.append(f"- Notes: {'; '.join(result.structure.notes[:3])}")
        lines.append("")
    return "\n".join(lines)


def run_stress_evaluation(
    *,
    fixtures: list[Path] | None = None,
    output_dir: Path | None = None,
    use_real_ai: bool = False,
    run_edit_checks: bool = True,
) -> dict:
    expectations = load_stress_expectations()
    fixture_paths = fixtures or ensure_all_stress_fixtures()
    timestamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
    run_dir = output_dir or (EVAL_ROOT / f"stress_{timestamp}")
    run_dir.mkdir(parents=True, exist_ok=True)
    work_dir = run_dir / "work"
    storage_dir = run_dir / "jobs"
    work_dir.mkdir(exist_ok=True)
    storage_dir.mkdir(exist_ok=True)

    settings = Settings.from_env(storage_dir=storage_dir)
    results: list[StressFixtureResult] = []

    mock_provider = pipeline_mock_provider()
    print("Running mock-provider stress pass...", flush=True)
    for fixture_path in fixture_paths:
        print(f"  mock: {fixture_path.name}", flush=True)
        result = evaluate_stress_fixture(
            fixture_path,
            expectations=expectations,
            provider=mock_provider,
            provider_mode="mock",
            settings=settings,
            work_dir=work_dir / "mock",
            storage_dir=storage_dir / "mock",
            run_edit_checks=run_edit_checks,
        )
        results.append(result)
        fixture_dir = run_dir / "mock" / fixture_path.stem
        fixture_dir.mkdir(parents=True, exist_ok=True)
        (fixture_dir / "issues.json").write_text(
            json.dumps([asdict(issue) for issue in result.issues], indent=2),
            encoding="utf-8",
        )

    real_ai_ran = False
    if use_real_ai:
        if not settings.provider_chain:
            print("Real AI skipped: no provider credentials configured.", file=sys.stderr)
        else:
            real_ai_ran = True
            real_provider = InstrumentedAIProvider(create_ai_provider(settings))
            print("Running real-AI stress pass...", flush=True)
            for fixture_path in fixture_paths:
                print(f"  real: {fixture_path.name}", flush=True)
                result = evaluate_stress_fixture(
                    fixture_path,
                    expectations=expectations,
                    provider=real_provider,
                    provider_mode="real",
                    settings=settings,
                    work_dir=work_dir / "real",
                    storage_dir=storage_dir / "real",
                    run_edit_checks=False,
                )
                results.append(result)
                fixture_dir = run_dir / "real" / fixture_path.stem
                fixture_dir.mkdir(parents=True, exist_ok=True)
                (fixture_dir / "issues.json").write_text(
                    json.dumps([asdict(issue) for issue in result.issues], indent=2),
                    encoding="utf-8",
                )

    mock_results = [result for result in results if result.provider_mode == "mock"]
    summary = {
        "timestamp": timestamp,
        "fixtures_evaluated": len(fixture_paths),
        "mock_completed": sum(1 for r in mock_results if r.status == "completed"),
        "mock_failed": sum(1 for r in mock_results if r.status == "failed"),
        "mock_error_issues": sum(
            1 for r in mock_results for issue in r.issues if issue.severity == "error"
        ),
        "real_ai_ran": real_ai_ran,
        "issue_counts": _issue_counts(results),
        "classified_issues": _classified_issues(results),
        "fixtures": [
            {
                "fixture": result.fixture,
                "provider_mode": result.provider_mode,
                "status": result.status,
                "elapsed_ms": round(result.elapsed_ms, 2),
                "detected_tasks": result.detected_tasks,
                "answerable_tasks": result.answerable_tasks,
                "answers_generated": result.answers_generated,
                "placement_strategies": result.placement_strategies,
                "missing_writes": result.missing_writes,
                "structure": asdict(result.structure),
                "edit": asdict(result.edit),
                "issues": [asdict(issue) for issue in result.issues],
            }
            for result in results
        ],
    }
    (run_dir / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    report = _render_inspection_report(mock_results, run_dir)
    (run_dir / "inspection_report.md").write_text(report, encoding="utf-8")
    return {"run_dir": str(run_dir), "summary": summary}


def main() -> None:
    parser = argparse.ArgumentParser(description="Run DocNA stress-test corpus evaluation")
    parser.add_argument("--fixture", help="Evaluate one stress fixture filename")
    parser.add_argument("--output-dir", type=Path, help="Override output directory")
    parser.add_argument("--real-ai", action="store_true", help="Also run real provider chain")
    parser.add_argument("--no-edit-checks", action="store_true", help="Skip editable review checks")
    args = parser.parse_args()

    fixtures = None
    if args.fixture:
        fixtures = [ensure_stress_fixture(args.fixture)]

    payload = run_stress_evaluation(
        fixtures=fixtures,
        output_dir=args.output_dir,
        use_real_ai=args.real_ai,
        run_edit_checks=not args.no_edit_checks,
    )
    summary = payload["summary"]
    print(json.dumps(summary, indent=2))
    print(f"\nReport written to: {payload['run_dir']}")


if __name__ == "__main__":
    main()
